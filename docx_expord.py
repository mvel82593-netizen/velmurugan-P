from io import BytesIO

from docx import Document
from docx.shared import Pt


def format_docx(title, content):
    """
    Create a DOCX document and return it as BytesIO.
    """

    doc = Document()

    # -----------------------------
    # Title
    # -----------------------------

    title_paragraph = doc.add_paragraph()

    title_run = title_paragraph.add_run(
        str(title)
    )

    title_run.bold = True
    title_run.font.size = Pt(18)

    # -----------------------------
    # Content
    # -----------------------------

    content_paragraph = doc.add_paragraph()

    content_run = content_paragraph.add_run(
        str(content)
    )

    content_run.font.size = Pt(11)

    # -----------------------------
    # Save to memory
    # -----------------------------

    output = BytesIO()

    doc.save(output)

    output.seek(0)

    return output